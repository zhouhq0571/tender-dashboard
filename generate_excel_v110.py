import json, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# 读取项目数据
with open('updated_projects.json') as f:
    projects = json.load(f)

wb = openpyxl.Workbook()

# 定义样式
header_fill = PatternFill(start_color='0A2540', end_color='0A2540', fill_type='solid')
header_font = Font(name='微软雅黑', size=10, bold=True, color='FFFFFF')
data_font = Font(name='微软雅黑', size=10)
center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
thin_border = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

# 标签颜色映射
tag_colors = {
    '🔥 ★★★ 强烈建议投标': 'FF6B6B',
    '⭐ ★★☆ 建议投标': 'FFD93D',
    '★☆☆ 可关注': 'C0C0C0',
    '👀 ★☆☆ 可关注': 'C0C0C0',
    '☆☆☆ 已截止': 'E0E0E0',
}

def create_sheet(ws, projects_filtered, sheet_name):
    ws.title = sheet_name
    headers = ['序号', '大区', '省份', '招标单位', '项目名称', '项目概述', '预算', '投标截止日期', '招标方式', '联系人', '业务标签', '投标建议', '信息来源', '链接']
    ws.append(headers)
    
    # 表头样式
    for col in range(1, len(headers)+1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
    
    # 数据行
    for i, p in enumerate(projects_filtered, 1):
        tags_str = ', '.join(p.get('tags', []))
        rec = p.get('rec', '')
        row = [
            i,
            p.get('region', ''),
            p.get('province', ''),
            p.get('company', ''),
            p.get('project', ''),
            p.get('overview', ''),
            p.get('budget', '-'),
            p.get('deadline', ''),
            p.get('method', ''),
            p.get('contact', '-'),
            tags_str,
            rec,
            p.get('source', ''),
            p.get('url', '')
        ]
        ws.append(row)
        
        # 数据行样式
        for col in range(1, len(headers)+1):
            cell = ws.cell(row=i+1, column=col)
            cell.font = data_font
            cell.alignment = center_align
            cell.border = thin_border
            # 投标建议列着色
            if col == 12 and rec in tag_colors:
                cell.fill = PatternFill(start_color=tag_colors[rec], end_color=tag_colors[rec], fill_type='solid')
    
    # 列宽
    ws.column_dimensions['A'].width = 6   # 序号
    ws.column_dimensions['B'].width = 8   # 大区
    ws.column_dimensions['C'].width = 8   # 省份
    ws.column_dimensions['D'].width = 14  # 招标单位
    ws.column_dimensions['E'].width = 30  # 项目名称
    ws.column_dimensions['F'].width = 40  # 项目概述
    ws.column_dimensions['G'].width = 12  # 预算
    ws.column_dimensions['H'].width = 16  # 截止日期
    ws.column_dimensions['I'].width = 12  # 招标方式
    ws.column_dimensions['J'].width = 30  # 联系人
    ws.column_dimensions['K'].width = 25  # 业务标签
    ws.column_dimensions['L'].width = 18  # 投标建议
    ws.column_dimensions['M'].width = 18  # 信息来源
    ws.column_dimensions['N'].width = 40  # 链接
    
    # 冻结窗格
    ws.freeze_panes = 'F2'
    
    # 行高
    ws.row_dimensions[1].height = 30
    for row in range(2, len(projects_filtered)+2):
        ws.row_dimensions[row].height = 40

# 全部有效招标
create_sheet(wb.active, projects, '全部有效招标')

# 银行行业
bank_projects = [p for p in projects if '信托' not in p.get('company', '') or '银行' in p.get('company', '')]
# 更准确的银行判断
bank_projects = [p for p in projects if not (
    (p.get('company', '').endswith('信托') and '银行' not in p.get('company', '')) 
    or p.get('company', '') == '中信登'
)]
wb.create_sheet()
create_sheet(wb.worksheets[1], bank_projects, '银行行业')

# 信托行业
trust_projects = [p for p in projects if (
    (p.get('company', '').endswith('信托') and '银行' not in p.get('company', '')) 
    or p.get('company', '') == '中信登'
)]
wb.create_sheet()
create_sheet(wb.worksheets[2], trust_projects, '信托行业')

# 强烈建议投标
strong = [p for p in projects if '强烈' in p.get('rec', '')]
wb.create_sheet()
create_sheet(wb.worksheets[3], strong, '强烈建议投标')

# 建议投标
suggest = [p for p in projects if '建议投标' in p.get('rec', '') and '强烈' not in p.get('rec', '')]
wb.create_sheet()
create_sheet(wb.worksheets[4], suggest, '建议投标')

# 保存
filename = '恒生银信招标资讯每日速递_2026年07月26日.xlsx'
wb.save(filename)
print(f"[Excel] 已生成: {filename}")
print(f"  - 全部有效招标: {len(projects)} 个")
print(f"  - 银行行业: {len(bank_projects)} 个")
print(f"  - 信托行业: {len(trust_projects)} 个")
print(f"  - 强烈建议投标: {len(strong)} 个")
print(f"  - 建议投标: {len(suggest)} 个")
