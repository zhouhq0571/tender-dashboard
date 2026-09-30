# -*- coding: utf-8 -*-
"""从看板 index.html 提取项目JSON，生成 2026-07-19 Excel 数据版（v105）"""
import re, json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

html = open('/Users/zhouhq/Documents/kimi/workspace/bidding-daily/index.html', encoding='utf-8').read()
m = re.search(r'"projects":\s*(\[)', html)
start = m.end() - 1
depth = 0; i = start
while i < len(html):
    if html[i] == '[': depth += 1
    elif html[i] == ']':
        depth -= 1
        if depth == 0: break
    i += 1
projects = json.loads(html[start:i+1])
print('项目总数:', len(projects))

all_projects = projects
bank_projects = [p for p in projects if '信托' not in p.get('company', '')]
trust_projects = [p for p in projects if '信托' in p.get('company', '')]
strong_projects = [p for p in projects if '强烈建议' in p.get('rec', '')]
suggest_projects = [p for p in projects if '建议投标' in p.get('rec', '') and '强烈建议' not in p.get('rec', '')]

wb = Workbook()
header_fill = PatternFill(start_color='0A2540', end_color='0A2540', fill_type='solid')
header_font = Font(color='FFFFFF', bold=True, size=10, name='微软雅黑')
data_font = Font(size=10, name='微软雅黑')
alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
thin_border = Border(
    left=Side(style='thin', color='CCCCCC'), right=Side(style='thin', color='CCCCCC'),
    top=Side(style='thin', color='CCCCCC'), bottom=Side(style='thin', color='CCCCCC'))

columns = [
    ('序号', 6), ('大区', 8), ('省份', 8), ('招标单位', 18),
    ('项目名称', 30), ('项目概况和招标范围', 50), ('预算金额', 15),
    ('投标截止日期', 18), ('招标方式', 12), ('项目咨询/联系方式', 25),
    ('业务标签', 25), ('投标建议', 15), ('信息来源', 25)]

def create_sheet(ws, sheet_projects, sheet_name):
    ws.title = sheet_name
    for col_idx, (col_name, col_width) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill; cell.font = header_font
        cell.alignment = alignment; cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = col_width
    for row_idx, p in enumerate(sheet_projects, 2):
        ws.row_dimensions[row_idx].height = 40
        values = [
            row_idx - 1, p.get('region', ''), p.get('province', ''), p.get('company', ''),
            p.get('project', ''), p.get('overview', ''), p.get('budget', ''),
            p.get('deadline', ''), p.get('method', ''), p.get('contact', ''),
            '、'.join(p.get('tags', [])), p.get('rec', ''), p.get('source', '')]
        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = data_font; cell.alignment = alignment; cell.border = thin_border
    ws.freeze_panes = 'F2'

create_sheet(wb.active, all_projects, '全部项目')
create_sheet(wb.create_sheet(), bank_projects, '银行')
create_sheet(wb.create_sheet(), trust_projects, '信托')
create_sheet(wb.create_sheet(), strong_projects, '强烈建议投标')
create_sheet(wb.create_sheet(), suggest_projects, '建议投标')

out = '/Users/zhouhq/Documents/kimi/workspace/恒生银信招标资讯每日速递_2026年7月19日.xlsx'
wb.save(out)
print('saved:', out)
print('全部:', len(all_projects), '银行:', len(bank_projects), '信托:', len(trust_projects),
      '强烈建议:', len(strong_projects), '建议投标:', len(suggest_projects))
