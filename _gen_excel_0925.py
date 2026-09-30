import json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# 读取数据
with open('/Users/zhouhq/Documents/kimi/workspace/bidding-daily/new_data.json', 'r', encoding='utf-8') as f:
    data = json.load(f)

projects = data['projects']

# 分类
all_projects = projects
bank_projects = [p for p in projects if '信托' not in p.get('company', '')]
trust_projects = [p for p in projects if '信托' in p.get('company', '')]
strong_projects = [p for p in projects if '强烈建议' in p.get('rec', '')]
suggest_projects = [p for p in projects if '建议投标' in p.get('rec', '')]

# 创建工作簿
wb = Workbook()

# 定义样式
header_fill = PatternFill(start_color='0A2540', end_color='0A2540', fill_type='solid')
header_font = Font(color='FFFFFF', bold=True, size=10, name='微软雅黑')
data_font = Font(size=10, name='微软雅黑')
alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
thin_border = Border(
    left=Side(style='thin', color='CCCCCC'),
    right=Side(style='thin', color='CCCCCC'),
    top=Side(style='thin', color='CCCCCC'),
    bottom=Side(style='thin', color='CCCCCC')
)

# 列定义
columns = [
    ('序号', 6), ('大区', 8), ('省份', 8), ('招标单位', 18),
    ('项目名称', 30), ('项目概况和招标范围', 50), ('预算金额', 15),
    ('投标截止日期', 18), ('招标方式', 12), ('项目咨询/联系方式', 25),
    ('业务标签', 25), ('投标建议', 15), ('信息来源', 25)
]

def create_sheet(ws, sheet_projects, sheet_name):
    ws.title = sheet_name
    
    # 写入表头
    for col_idx, (col_name, col_width) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = alignment
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = col_width
    
    # 写入数据
    for row_idx, p in enumerate(sheet_projects, 2):
        ws.row_dimensions[row_idx].height = 40
        
        values = [
            row_idx - 1,  # 序号
            p.get('region', ''),
            p.get('province', ''),
            p.get('company', ''),
            p.get('project', ''),
            p.get('overview', ''),
            p.get('budget', ''),
            p.get('deadline', ''),
            p.get('method', ''),
            p.get('contact', ''),
            '、'.join(p.get('tags', [])),
            p.get('rec', ''),
            p.get('source', '')
        ]
        
        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = data_font
            cell.alignment = alignment
            cell.border = thin_border
    
    # 冻结窗格
    ws.freeze_panes = 'F2'

# 创建各Sheet
create_sheet(wb.active, all_projects, f'全部有效招标（{len(all_projects)}）')

wb.create_sheet()
create_sheet(wb.worksheets[1], bank_projects, f'银行行业（{len(bank_projects)}）')

wb.create_sheet()
create_sheet(wb.worksheets[2], trust_projects, f'信托行业（{len(trust_projects)}）')

wb.create_sheet()
create_sheet(wb.worksheets[3], strong_projects, f'强烈建议投标（{len(strong_projects)}）')

wb.create_sheet()
create_sheet(wb.worksheets[4], suggest_projects, f'建议投标（{len(suggest_projects)}）')

# 保存Excel
excel_path = '/Users/zhouhq/Documents/kimi/workspace/bidding-daily/恒生银信招标资讯每日速递_2026年9月25日.xlsx'
wb.save(excel_path)

print(f"Excel文件已保存: {excel_path}")
print(f"Sheet1: 全部有效招标（{len(all_projects)}）")
print(f"Sheet2: 银行行业（{len(bank_projects)}）")
print(f"Sheet3: 信托行业（{len(trust_projects)}）")
print(f"Sheet4: 强烈建议投标（{len(strong_projects)}）")
print(f"Sheet5: 建议投标（{len(suggest_projects)}）")
